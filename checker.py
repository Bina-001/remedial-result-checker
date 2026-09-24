from flask import Flask, render_template, request, jsonify, Response
import requests
import csv
import io
import openpyxl
import json
import os
import time
from datetime import datetime

app = Flask(__name__)

MINISTRY_API_URL = "https://eap.ethernet.edu.et/api/remedial-exam-result"
MINISTRY_API_KEY = os.getenv("MINISTRY_API_KEY", "")

DATA_FILE = "results.json"
HISTORY_FILE = "history.json"

MAX_RETRIES = 3
RETRY_DELAYS = [5, 15, 30]

uploaded_students = []
checked_results = []

history_data = []
current_history_id = ""


def save_results():
    data = {
        "uploaded_students": uploaded_students,
        "checked_results": checked_results
    }

    with open(DATA_FILE, "w", encoding="utf-8") as file:
        json.dump(data, file, ensure_ascii=False, indent=2)


def save_history():
    data = {
        "histories": history_data
    }

    with open(HISTORY_FILE, "w", encoding="utf-8") as file:
        json.dump(data, file, ensure_ascii=False, indent=2)


def load_data():
    global uploaded_students
    global checked_results
    global history_data
    global current_history_id

    if os.path.exists(DATA_FILE):
        try:
            with open(DATA_FILE, "r", encoding="utf-8") as file:
                data = json.load(file)

            uploaded_students = data.get(
                "uploaded_students",
                []
            )

            checked_results = data.get(
                "checked_results",
                [])

        except Exception:
            uploaded_students = []
            checked_results = []

    if os.path.exists(HISTORY_FILE):
        try:
            with open(HISTORY_FILE, "r", encoding="utf-8") as file:
                data = json.load(file)

            history_data = data.get(
                "histories",
                []
            )

        except Exception:
            history_data = []

    if not history_data and (
        uploaded_students or checked_results
    ):
        history_id = datetime.now().strftime(
            "%Y%m%d_%H%M%S"
        )

        history_data.append({
            "id": history_id,
            "name": "Previous Results",
            "created_at": datetime.now().isoformat(),
            "updated_at": datetime.now().isoformat(),
            "students": uploaded_students,
            "results": checked_results
        })

        current_history_id = history_id
        save_history()

    elif history_data:
        current_history_id = history_data[-1]["id"]


def get_current_history():
    for history in history_data:
        if history.get("id") == current_history_id:
            return history

    return None


def classify_result(stream, total):
    stream_text = str(
        stream or ""
    ).strip().lower()

    total_text = str(
        total or ""
    ).strip()

    if total_text.lower() == "disqualified":
        return "DISQUALIFIED"

    try:
        numeric_total = float(total_text)
    except (ValueError, TypeError):
        return "OTHER"

    if "social" in stream_text:
        if numeric_total >= 200:
            return "PASS"

        return "FAIL"

    if "natural" in stream_text:
        if numeric_total >= 250:
            return "PASS"

        return "FAIL"

    return "OTHER"


def save_or_update_result(result):
    global checked_results

    username = str(
        result.get("username", "")
    ).strip()

    if not username:
        return

    replaced = False

    for index, existing in enumerate(checked_results):
        existing_username = str(
            existing.get("username", "")
        ).strip()

        if existing_username == username:
            checked_results[index] = result
            replaced = True
            break

    if not replaced:
        checked_results.append(result)

    save_results()

    history = get_current_history()

    if history is not None:

        history_results = history.setdefault(
            "results",
            []
        )

        replaced_history = False

        for index, existing in enumerate(history_results):
            existing_username = str(
                existing.get("username", "")
            ).strip()

            if existing_username == username:
                history_results[index] = result
                replaced_history = True
                break

        if not replaced_history:
            history_results.append(result)

        history["updated_at"] = datetime.now().isoformat()

        save_history()


def create_history(students):
    global history_data
    global current_history_id
    global uploaded_students
    global checked_results

    history_id = datetime.now().strftime(
        "%Y%m%d_%H%M%S_%f"
    )

    history_name = (
        "Roster "
        + datetime.now().strftime(
            "%Y-%m-%d %H:%M"
        )
    )

    history = {
        "id": history_id,
        "name": history_name,
        "created_at": datetime.now().isoformat(),
        "updated_at": datetime.now().isoformat(),
        "students": students,
        "results": []
    }

    history_data.append(history)

    current_history_id = history_id

    uploaded_students = students
    checked_results = []

    save_results()
    save_history()


@app.route("/")
def home():
    return render_template("index.html")


@app.route("/health")
def health():
    return jsonify({
        "success": True,
        "message": "Remedial Result Checker is running"
    })


@app.route("/upload", methods=["POST"])
def upload_file():
    file = request.files.get("file")

    if not file:
        return jsonify({
            "success": False,
            "message": "No file uploaded"
        }), 400

    filename = (
        file.filename or ""
    ).lower()

    try:
        students = []

        if filename.endswith(".csv"):

            content = file.read().decode(
                "utf-8-sig"
            )

            reader = csv.DictReader(
                io.StringIO(content)
            )

            if not reader.fieldnames:
                return jsonify({
                    "success": False,
                    "message": "CSV has no header row"
                }), 400

            username_column = None

            for column in reader.fieldnames:
                if (
                    column is not None
                    and column.strip().lower()
                    == "username"
                ):
                    username_column = column
                    break

            if username_column is None:
                return jsonify({
                    "success": False,
                    "message":
                        "Username column was not found"
                }), 400

            for row_number, row in enumerate(
                reader,
                start=2
            ):
                value = row.get(
                    username_column
                )

                username = (
                    str(value).strip()
                    if value is not None
                    else ""
                )

                students.append({
                    "row_number": row_number,
                    "username": username,
                    "status":
                        "READY"
                        if username
                        else "MISS"
                })

        elif filename.endswith(".xlsx"):

            workbook = openpyxl.load_workbook(
                file,
                read_only=True,
                data_only=True
            )

            sheet = workbook.active

            rows = sheet.iter_rows(
                values_only=True
            )

            try:
                headers = next(rows)
            except StopIteration:

                workbook.close()

                return jsonify({
                    "success": False,
                    "message":
                        "Excel file is empty"
                }), 400

            username_index = None

            for index, header in enumerate(headers):
                if (
                    header is not None
                    and str(header).strip().lower()
                    == "username"
                ):
                    username_index = index
                    break

            if username_index is None:

                workbook.close()

                return jsonify({
                    "success": False,
                    "message":
                        "Username column was not found"
                }), 400

            for row_number, row in enumerate(
                rows,
                start=2
            ):

                value = (
                    row[username_index]
                    if username_index < len(row)
                    else None
                )

                username = (
                    str(value).strip()
                    if value is not None
                    else ""
                )

                students.append({
                    "row_number": row_number,
                    "username": username,
                    "status":
                        "READY"
                        if username
                        else "MISS"
                })

            workbook.close()

        else:

            return jsonify({
                "success": False,
                "message":
                    "Please upload a CSV or XLSX file"
            }), 400

        if not students:

            return jsonify({
                "success": False,
                "message":
                    "No student rows were found"
            }), 400

        create_history(students)

        ready_count = sum(
            1
            for student in students
            if student["status"] == "READY"
        )

        missing_count = (
            len(students) - ready_count
        )

        first_username = ""

        for student in students:
            if student["username"]:
                first_username = student["username"]
                break

        return jsonify({
            "success": True,
            "message":
                "New history created successfully",
            "history_id":
                current_history_id,
            "total_students":
                len(students),
            "ready_students":
                ready_count,
            "missing_usernames":
                missing_count,
            "test_username":
                first_username
        })

    except Exception as error:

        return jsonify({
            "success": False,
            "message":
                f"File processing error: {error}"
        }), 500


@app.route("/students")
def get_students():

    return jsonify({
        "success": True,
        "history_id":
            current_history_id,
        "total_students":
            len(uploaded_students),
        "students":
            uploaded_students
    })


@app.route("/results")
def get_results():

    return jsonify({
        "success": True,
        "history_id":
            current_history_id,
        "results":
            checked_results
    })


@app.route("/history")
def get_history():

    items = []

    for history in reversed(history_data):

        students = history.get(
            "students",
            []
        )

        results = history.get(
            "results",
            []
        )

        items.append({
            "id":
                history.get("id", ""),
            "name":
                history.get("name", ""),
            "created_at":
                history.get("created_at", ""),
            "updated_at":
                history.get("updated_at", ""),
            "total_students":
                len(students),
            "checked":
                len(results),
            "is_current":
                history.get("id")
                == current_history_id
        })

    return jsonify({
        "success": True,
        "histories": items
    })


@app.route("/history/<history_id>")
def open_history(history_id):

    global uploaded_students
    global checked_results
    global current_history_id

    for history in history_data:

        if history.get("id") == history_id:

            current_history_id = history_id

            uploaded_students = history.get(
                "students",
                []
            )

            checked_results = history.get(
                "results",
                []
            )

            save_results()

            return jsonify({
                "success": True,
                "message":
                    "History opened",
                "history_id":
                    history_id,
                "students":
                    uploaded_students,
                "results":
                    checked_results
            })

    return jsonify({
        "success": False,
        "message":
            "History was not found"
    }), 404


@app.route("/search")
def search_student():

    username = str(
        request.args.get(
            "username",
            ""
        )
    ).strip()

    if not username:

        return jsonify({
            "success": False,
            "message":
                "Username is required"
        }), 400

    for result in checked_results:

        result_username = str(
            result.get(
                "username",
                ""
            )
        ).strip()

        if result_username == username:

            return jsonify({
                "success": True,
                "found": True,
                "checked": True,
                "result": result,
                "history_id":
                    current_history_id
            })

    for student in uploaded_students:

        student_username = str(
            student.get(
                "username",
                ""
            )
        ).strip()

        if student_username == username:

            return jsonify({
                "success": True,
                "found": True,
                "checked": False,
                "student": student,
                "message":
                    "Student has not been checked yet",
                "history_id":
                    current_history_id
            })

    return jsonify({
        "success": True,
        "found": False,
        "checked": False,
        "message":
            "Student was not found in this history"
    })


@app.route("/check", methods=["POST"])
def check_result():

    data = request.get_json(
        silent=True
    ) or {}

    username = str(
        data.get("username") or ""
    ).strip()

    if not username:

        return jsonify({
            "success": False,
            "status": "ERROR",
            "score": "",
            "username": "",
            "message":
                "Username is missing"
        })

    for result in checked_results:

        existing_username = str(
            result.get(
                "username",
                ""
            )
        ).strip()

        if existing_username == username:

            return jsonify({
                **result,
                "cached": True,
                "message":
                    "Result loaded from history"
            })

    for attempt in range(
        MAX_RETRIES + 1
    ):

        try:

            response = requests.post(
                MINISTRY_API_URL,
                headers={
                    "X-API-Key":
                        MINISTRY_API_KEY,
                    "Content-Type":
                        "application/json",
                    "Accept":
                        "application/json"
                },
                json={
                    "username":
                        username
                },
                timeout=20
            )

            try:
                result_data = response.json()

            except ValueError:

                result = {
                    "success": False,
                    "status": "ERROR",
                    "score": "",
                    "username":
                        username,
                    "message":
                        "Ministry returned an invalid response"
                }

                save_or_update_result(result)

                return jsonify(result)

            if response.status_code == 429:

                if attempt < MAX_RETRIES:

                    time.sleep(
                        RETRY_DELAYS[attempt]
                    )

                    continue

                result = {
                    "success": False,
                    "status": "ERROR",
                    "score": "",
                    "username":
                        username,
                    "message":
                        "Ministry rate limit reached after retries"
                }

                save_or_update_result(result)

                return jsonify(result)

            if response.status_code != 200:

                result = {
                    "success": False,
                    "status": "ERROR",
                    "score": "",
                    "username":
                        username,
                    "message":
                        result_data.get(
                            "message",
                            f"Ministry request failed: HTTP {response.status_code}"
                        )
                }

                save_or_update_result(result)

                return jsonify(result)

            if result_data.get(
                "status"
            ) != "success":

                result = {
                    "success": False,
                    "status": "ERROR",
                    "score": "",
                    "username":
                        username,
                    "message":
                        result_data.get(
                            "message",
                            "Ministry did not return a result"
                        )
                }

                save_or_update_result(result)

                return jsonify(result)

            ministry_data = result_data.get(
                "data",
                {}
            )

            total = ministry_data.get(
                "total"
            )

            stream = ministry_data.get(
                "stream_name",
                ""
            )

            status = classify_result(
                stream,
                total
            )

            result = {
                "success": True,
                "status": status,
                "official_total": total,
                "score": total,
                "username":
                    ministry_data.get(
                        "username",
                        username
                    ),
                "message":
                    result_data.get(
                        "message",
                        "Result retrieved from Ministry"
                    ),
                "full_name":
                    ministry_data.get(
                        "full_name",
                        ""
                    ),
                "year":
                    ministry_data.get(
                        "short_year",
                        ""
                    ),
                "gender":
                    ministry_data.get(
                        "gender",
                        ""
                    ),
                "institution":
                    ministry_data.get(
                        "institution_name",
                        ""
                    ),
                "stream":
                    stream,
                "subjects":
                    ministry_data.get(
                        "subjects",
                        []
                    ),
                "checked_at":
                    datetime.now().isoformat()
            }

            save_or_update_result(result)

            return jsonify(result)

        except requests.Timeout:

            if attempt < MAX_RETRIES:

                time.sleep(
                    RETRY_DELAYS[attempt]
                )

                continue

            result = {
                "success": False,
                "status": "ERROR",
                "score": "",
                "username":
                    username,
                "message":
                    "Ministry request timed out after retries"
            }

            save_or_update_result(result)

            return jsonify(result)

        except requests.RequestException as error:

            result = {
                "success": False,
                "status": "ERROR",
                "score": "",
                "username":
                    username,
                "message":
                    f"Connection error: {error}"
            }

            save_or_update_result(result)

            return jsonify(result)

    return jsonify({
        "success": False,
        "status": "ERROR",
        "score": "",
        "username":
            username,
        "message":
            "Request failed"
    })


@app.route("/summary")
def summary():

    total = len(
        uploaded_students
    )

    ready = sum(
        1
        for student in uploaded_students
        if str(
            student.get(
                "username",
                ""
            )
        ).strip()
    )

    completed_usernames = set()

    for result in checked_results:

        username = str(
            result.get(
                "username",
                ""
            )
        ).strip()

        status = str(
            result.get(
                "status",
                ""
            )
        ).strip().upper()

        if username and status != "ERROR":

            completed_usernames.add(
                username
            )

    completed = len(
        completed_usernames
    )

    remaining = max(
        ready - completed,
        0
    )

    pass_count = 0
    fail_count = 0
    disqualified_count = 0
    error_count = 0
    other_count = 0

    for result in checked_results:

        status = str(
            result.get(
                "status",
                ""
            )
        ).strip().upper()

        if status == "PASS":

            pass_count += 1

        elif status == "FAIL":

            fail_count += 1

        elif status == "DISQUALIFIED":

            disqualified_count += 1

        elif status == "ERROR":

            error_count += 1

        else:

            other_count += 1

    return jsonify({
        "success": True,
        "history_id":
            current_history_id,
        "total":
            total,
        "ready":
            ready,
        "missing_usernames":
            total - ready,
        "completed":
            completed,
        "remaining":
            remaining,
        "pass":
            pass_count,
        "fail":
            fail_count,
        "disqualified":
            disqualified_count,
        "error":
            error_count,
        "other":
            other_count
    })


@app.route("/resume")
def resume():

    completed_usernames = set()

    for result in checked_results:

        username = str(
            result.get(
                "username",
                ""
            )
        ).strip()

        status = str(
            result.get(
                "status",
                ""
            )
        ).strip().upper()

        if username and status != "ERROR":

            completed_usernames.add(
                username
            )

    remaining = []

    for student in uploaded_students:

        username = str(
            student.get(
                "username",
                ""
            )
        ).strip()

        if (
            username
            and username not in
            completed_usernames
        ):

            remaining.append(
                student
            )

    return jsonify({
        "success": True,
        "remaining":
            len(remaining),
        "students":
            remaining
    })


@app.route("/export/csv")
def export_csv():

    output = io.StringIO()

    writer = csv.writer(
        output
    )

    writer.writerow([
        "No",
        "Username",
        "Name",
        "Institution",
        "Stream",
        "Status",
        "Score",
        "Message"
    ])

    for number, result in enumerate(
        checked_results,
        start=1
    ):

        writer.writerow([
            number,
            result.get(
                "username",
                ""
            ),
            result.get(
                "full_name",
                ""
            ),
            result.get(
                "institution",
                ""
            ),
            result.get(
                "stream",
                ""
            ),
            result.get(
                "status",
                ""
            ),
            result.get(
                "score",
                ""
            ),
            result.get(
                "message",
                ""
            )
        ])

    return Response(
        output.getvalue(),
        mimetype="text/csv",
        headers={
            "Content-Disposition":
                "attachment; filename=remedial_results.csv"
        }
    )


@app.route("/export/xlsx")
def export_xlsx():

    workbook = openpyxl.Workbook()

    results_sheet = workbook.active
    results_sheet.title = "Results"

    results_sheet.append([
        "No",
        "Username",
        "Name",
        "Institution",
        "Stream",
        "Status",
        "Total",
        "Message"
    ])

    for number, result in enumerate(
        checked_results,
        start=1
    ):

        results_sheet.append([
            number,
            result.get(
                "username",
                ""
            ),
            result.get(
                "full_name",
                ""
            ),
            result.get(
                "institution",
                ""
            ),
            result.get(
                "stream",
                ""
            ),
            result.get(
                "status",
                ""
            ),
            result.get(
                "score",
                ""
            ),
            result.get(
                "message",
                ""
            )
        ])

    subjects_sheet = workbook.create_sheet(
        title="Subjects"
    )

    subjects_sheet.append([
        "Username",
        "Name",
        "Subject",
        "Score"
    ])

    for result in checked_results:

        subjects = result.get(
            "subjects",
            []
        )

        if isinstance(subjects, list):

            for subject in subjects:

                subjects_sheet.append([
                    result.get(
                        "username",
                        ""
                    ),
                    result.get(
                        "full_name",
                        ""
                    ),
                    subject.get(
                        "name",
                        ""
                    ),
                    subject.get(
                        "score",
                        ""
                    )
                ])

    output = io.BytesIO()

    workbook.save(output)

    output.seek(0)

    return Response(
        output.getvalue(),
        mimetype=
            "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={
            "Content-Disposition":
                "attachment; filename=remedial_results.xlsx"
        }
    )


load_data()


if __name__ == "__main__":

    app.run(
        host="0.0.0.0",
        port=5000,
        debug=True
    )
